"""Report source file checksums, spreadsheet metadata and archive inventories."""
import collections
import hashlib
import json
from pathlib import Path
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

def main():
    for p in sorted((ROOT / "data").iterdir()):
        print("SOURCE", p.name, p.stat().st_size, hashlib.sha256(p.read_bytes()).hexdigest())
        if p.suffix == ".xlsx":
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            for ws in wb:
                print("SHEET", ws.title, ws.max_row, ws.max_column)
                rows = list(ws.iter_rows(values_only=True))
                print("CATEGORIES", collections.Counter(r[0] for r in rows[1:]))
                print("MISSING_METADATA", [r for r in rows[1:] if r[5] in [743474,743475]])
                for row in rows[:5]:
                    print(json.dumps(row, default=str))
            wb.close()
        if p.suffix == ".zip":
            with zipfile.ZipFile(p) as z:
                members = [i for i in z.infolist() if not i.is_dir()]
                print("ZIP", len(members), sum(i.file_size for i in members))
                headers=collections.Counter()
                for member in members:
                    with z.open(member) as stream:
                        headers[stream.readline(4096).decode("utf-8",errors="replace").strip()]+=1
                print("ALL_HEADERS", json.dumps(dict(headers)))
                print("EXTENSIONS", collections.Counter(Path(i.filename).suffix for i in members))
                print("NAMES", json.dumps([i.filename for i in members[:25]]))
                representatives = {}
                for i in members:
                    key = ("trades" if i.filename.endswith("-trades.csv") else "daily", i.filename.split("/")[0])
                    representatives.setdefault(key, i)
                for i in list(representatives.values())[:4]:
                    with z.open(i) as stream:
                        print("SAMPLE", i.filename, repr(stream.read(1600)))
    p = ROOT / "data/legacy-datastore.bin"
    with p.open("rb") as f:
        print("DATASTORE_HEADER", repr(f.read(128)))
    print("DATASTORE_SHA256", hashlib.sha256(p.read_bytes()).hexdigest())

if __name__ == "__main__":
    main()
